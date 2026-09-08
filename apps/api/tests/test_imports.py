"""Ввоз документов.

Проверяется то, что нельзя починить задним числом: отказ до работы, пределы
распаковки, путь наружу каталога и запись причины отказа в само задание.

Разбор документов сюда не дублируется — он проверяется у своего модуля
(`test_document_text.py`) и у сервиса преобразования. Здесь проверяется ввоз:
что приходит, что отвергается и что остаётся в базе.
"""

from __future__ import annotations

import base64
import io
import json
import re
import uuid
import zipfile
from pathlib import Path

import httpx
import pytest
from sqlalchemy import func, insert, select
from sqlalchemy.ext.asyncio import AsyncSession

from tessera_api.domain.errors import AppError
from tessera_api.infrastructure.content import (
    PDF_FAILURES,
    ContentClient,
    _failure_code,
)
from tessera_api.infrastructure.models import (
    Attachment,
    BaseProperty,
    BaseRow,
    BaseView,
    Comment,
    FileTask,
    Label,
    Page,
    PageLabel,
    PageVerification,
    PageVerifier,
    Share,
    SpaceMember,
    User,
)
from tessera_api.services import imports as module
from tessera_api.services.imports import (
    MAX_ENTRIES,
    SINGLE_FILE_EXTENSIONS,
    SOURCES,
    STATUS_FAILED,
    STATUS_PROCESSING,
    STATUS_SUCCESS,
    ArchiveEntry,
    ImportService,
    _listing,
    _parent_for,
    _unwrapped,
    _without_notion_twins,
    assert_supported,
    drop_title_heading,
    read_table,
    safe_entries,
    table_to_html,
    title_from_html,
    title_from_markdown,
)
from tests.conftest import RealtimeDouble, needs_database

DOC = {"type": "doc", "content": [{"type": "paragraph"}]}


def content_client(
    seen: list | None = None, *, pdf: str | None = None, doc: dict | None = None
) -> ContentClient:
    """Сервис преобразования, отвечающий пустым документом.

    Настоящий здесь не нужен: схема узлов живёт в нём, и проверять её второй
    раз значило бы завести второе описание того же.

    Разбор PDF по умолчанию отвечает отказом «нет текстового слоя»: разбирать
    PDF умеет только настоящий сервис, а проверкам нужен предсказуемый ответ.
    Разметка задаётся доводом `pdf` там, где проверяется успешный ввоз.

    Довод `doc` задаёт разобранный документ там, где проверяется обработка
    разобранного, а не сам разбор.
    """

    def handler(request: httpx.Request) -> httpx.Response:
        if seen is not None:
            seen.append(json.loads(request.content))
        if request.url.path.endswith("/pdf-to-html"):
            if pdf is None:
                return httpx.Response(400, json={"error": "pdf_no_text_layer"})
            return httpx.Response(200, json={"html": pdf})
        return httpx.Response(200, json={"content": doc or DOC})

    return ContentClient("http://collab:3001", transport=httpx.MockTransport(handler))


def archive(files: dict[str, bytes | str], *, compress: bool = True) -> bytes:
    buffer = io.BytesIO()
    mode = zipfile.ZIP_DEFLATED if compress else zipfile.ZIP_STORED
    with zipfile.ZipFile(buffer, "w", mode) as bundle:
        for name, body in files.items():
            bundle.writestr(name, body)
    return buffer.getvalue()


class TestSupportedFormats:
    def test_an_unknown_extension_is_refused(self) -> None:
        with pytest.raises(AppError) as error:
            assert_supported("выгрузка.rtf")
        assert error.value.code == "error.import.unsupported_format"

    def test_the_refusal_names_what_is_accepted(self) -> None:
        """Выгрузка из чужого редактора предлагает семь форматов.

        Без перечисления человек не понимает, что именно менять при выгрузке, и
        приходит спрашивать.
        """
        with pytest.raises(AppError) as error:
            assert_supported("документ.epub")
        assert ".md" in (error.value.extra or {}).get("params", {}).get("allowed", "")

    def test_an_open_document_is_accepted(self) -> None:
        """ODT принимается наравне с DOCX: тот же текст, тот же путь ввоза."""
        assert assert_supported("документ.odt") == ".odt"

    def test_every_declared_format_passes(self) -> None:
        for suffix in SINGLE_FILE_EXTENSIONS:
            assert assert_supported(f"файл{suffix}") == suffix

    def test_the_case_of_the_extension_does_not_matter(self) -> None:
        """Выгрузка из редактора под другой системой даёт `.MD` и `.PDF`."""
        assert assert_supported("Документ.PDF") == ".pdf"

    def test_a_name_without_an_extension_is_refused(self) -> None:
        with pytest.raises(AppError):
            assert_supported("документ")


class TestTitles:
    def test_the_document_title_wins(self) -> None:
        html = "<html><head><title>Из заголовка</title></head><body><h1>Из текста</h1></body>"
        assert title_from_html(html, "файл.html") == "Из заголовка"

    def test_the_first_heading_is_next(self) -> None:
        assert title_from_html("<body><h1>Из текста</h1></body>", "файл.html") == "Из текста"

    def test_markup_inside_the_heading_is_dropped(self) -> None:
        assert title_from_html("<h1>Часть <em>вторая</em></h1>", "файл.html") == "Часть вторая"

    def test_an_empty_title_does_not_win(self) -> None:
        """Пустой заголовок документа — не название, а его отсутствие."""
        assert title_from_html("<title>  </title><h1>Из текста</h1>", "ф.html") == "Из текста"

    def test_the_file_name_is_the_last_resort(self) -> None:
        assert title_from_html("<p>Без заголовков</p>", "Отчёт за год.html") == "Отчёт за год"

    def test_markdown_takes_the_leading_heading(self) -> None:
        assert title_from_markdown("# Название\n\nтекст", "файл.md") == "Название"

    def test_a_heading_after_text_is_not_the_title(self) -> None:
        """Заголовок в середине документа — это раздел, а не название."""
        assert title_from_markdown("Вступление\n\n# Раздел", "Файл.md") == "Файл"

    def test_a_leading_empty_line_does_not_hide_the_heading(self) -> None:
        assert title_from_markdown("\n\n# Название", "файл.md") == "Название"


class TestTitleHeading:
    """Снятие заголовка, повторяющего название.

    Вывоз пишет название первым заголовком тела, ввоз берёт название оттуда же.
    Без снятия оборот «вывоз — ввоз» добавлял странице по заголовку за раз.
    """

    def _doc(self, *nodes) -> dict:
        return {"type": "doc", "content": list(nodes)}

    def _heading(self, text: str, level: int = 1) -> dict:
        return {
            "type": "heading",
            "attrs": {"level": level},
            "content": [{"type": "text", "text": text}],
        }

    def test_the_repeated_heading_goes(self) -> None:
        body = self._doc(self._heading("Отчёт"), {"type": "paragraph"})
        assert drop_title_heading(body, "Отчёт")["content"] == [{"type": "paragraph"}]

    def test_another_heading_stays(self) -> None:
        # Заголовок, не совпадающий с названием, — часть документа.
        body = self._doc(self._heading("Раздел"), {"type": "paragraph"})
        assert len(drop_title_heading(body, "Отчёт")["content"]) == 2

    def test_a_lower_level_heading_stays(self) -> None:
        body = self._doc(self._heading("Отчёт", level=2), {"type": "paragraph"})
        assert len(drop_title_heading(body, "Отчёт")["content"]) == 2

    def test_a_document_of_one_heading_keeps_a_paragraph(self) -> None:
        # Пустой документ редактор не принимает.
        body = self._doc(self._heading("Отчёт"))
        assert drop_title_heading(body, "Отчёт")["content"] == [{"type": "paragraph"}]

    def test_an_empty_document_is_left_alone(self) -> None:
        assert drop_title_heading({"type": "doc", "content": []}, "Отчёт") == {
            "type": "doc",
            "content": [],
        }

    def test_the_case_and_the_spacing_do_not_matter(self) -> None:
        """Выгрузка пишет название в теле не тем же написанием, что в оглавлении.

        При точном сравнении заголовок оставался бы задвоенным ровно там, где
        он и задваивается.
        """
        body = self._doc(self._heading("ОТЧЁТ   за  год"), {"type": "paragraph"})
        assert drop_title_heading(body, "Отчёт за год")["content"] == [{"type": "paragraph"}]

    def test_an_empty_heading_stays(self) -> None:
        """Пустой заголовок не совпадает ни с каким названием.

        Снятый, он был бы правкой документа там, где о повторе названия речи
        не было.
        """
        body = self._doc(self._heading(""), {"type": "paragraph"})
        assert len(drop_title_heading(body, "Отчёт")["content"]) == 2

    def test_a_page_without_a_title_keeps_its_heading(self) -> None:
        body = self._doc(self._heading("Отчёт"), {"type": "paragraph"})
        assert len(drop_title_heading(body, "")["content"]) == 2


class TestArchiveSafety:
    def test_a_broken_archive_is_a_refusal(self) -> None:
        with pytest.raises(AppError) as error:
            safe_entries(b"\x00\x01\x02" + "это не архив".encode())
        assert error.value.code == "error.import.broken_archive"

    def test_a_bomb_is_stopped_by_the_budget(self) -> None:
        """Архив, помещающийся в письмо, не должен разворачиваться в терабайт."""
        bomb = archive({"огромный.md": b"\x00" * (300 * 1024 * 1024)})
        assert len(bomb) < 1024 * 1024

        with pytest.raises(AppError) as error:
            safe_entries(bomb)
        assert error.value.code == "error.import.archive_too_large"

    def test_an_honest_archive_fits_the_budget(self) -> None:
        """Обычная выгрузка сжимается вчетверо и обязана проходить."""
        text = ("Обычный текст страницы. " * 4000).encode()
        entries = safe_entries(archive({"страница.md": text}))
        assert [one.path for one in entries] == ["страница.md"]

    def test_a_path_leaving_the_folder_is_dropped(self) -> None:
        entries = safe_entries(
            archive({"../снаружи.md": "чужое", "внутри.md": "своё"})
        )
        assert [one.path for one in entries] == ["внутри.md"]

    def test_a_deep_path_leaving_the_folder_is_dropped_too(self) -> None:
        """Выход наружу бывает не только первым знаком пути."""
        entries = safe_entries(archive({"папка/../../снаружи.md": "чужое"}))
        assert entries == []

    def test_an_absolute_path_is_dropped(self) -> None:
        entries = safe_entries(archive({"/etc/passwd": "чужое", "своя.md": "своё"}))
        assert [one.path for one in entries] == ["своя.md"]

    def test_the_junk_of_the_archiver_is_skipped(self) -> None:
        """Без отсева служебного каталога каждая страница ввозится дважды."""
        entries = safe_entries(
            archive(
                {
                    "__MACOSX/._страница.md": "копия",
                    ".DS_Store": "мусор",
                    "страница.md": "текст",
                }
            )
        )
        assert [one.path for one in entries] == ["страница.md"]

    def test_the_number_of_entries_is_capped(self, monkeypatch: pytest.MonkeyPatch) -> None:
        """Архив из миллиона пустых файлов опасен обходом, а не объёмом."""
        monkeypatch.setattr(module, "MAX_ENTRIES", 2)
        entries = safe_entries(archive({f"файл{index}.md": "текст" for index in range(5)}))
        assert len(entries) == 2

    def test_the_cap_is_high_enough_for_a_real_wiki(self) -> None:
        assert MAX_ENTRIES >= 100_000

    def test_folders_are_not_entries(self) -> None:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as bundle:
            bundle.writestr("папка/", "")
            bundle.writestr("папка/страница.md", "текст")
        assert [one.path for one in safe_entries(buffer.getvalue())] == ["папка/страница.md"]


class TestListing:
    def test_no_listing_is_an_ordinary_case(self) -> None:
        found = _listing(safe_entries(archive({"страница.md": "текст"})))
        assert found.known == {}

    def test_icons_and_order_come_from_the_listing(self) -> None:
        meta = {
            "source": "tessera",
            "pages": {
                "Вторая.md": {"icon": "📗", "position": "hz"},
                "Первая.md": {"icon": "📘", "position": "ha"},
            },
        }
        found = _listing(
            safe_entries(
                archive(
                    {
                        "tessera-metadata.json": json.dumps(meta),
                        "Первая.md": "# Первая",
                        "Вторая.md": "# Вторая",
                    }
                )
            )
        )
        assert found.about("Первая.md")["icon"] == "📘"
        assert found.rank("Первая.md") < found.rank("Вторая.md")

    def test_an_unknown_entry_goes_last(self) -> None:
        """Оглавление не задаёт места чужому файлу, и выдумывать его нельзя."""
        meta = {"pages": {"Своя.md": {"position": "ha"}}}
        found = _listing(
            safe_entries(
                archive({"tessera-metadata.json": json.dumps(meta), "Своя.md": "текст"})
            )
        )
        assert found.rank("Чужая.md") > found.rank("Своя.md")

    def test_a_listing_one_level_down_still_matches(self) -> None:
        """Архив нередко собирают из каталога, и всё внутри сдвинуто."""
        meta = {"pages": {"Первая.md": {"icon": "📘"}}}
        found = _listing(
            safe_entries(
                archive(
                    {
                        "выгрузка/tessera-metadata.json": json.dumps(meta),
                        "выгрузка/Первая.md": "текст",
                    }
                )
            )
        )
        assert found.about("выгрузка/Первая.md")["icon"] == "📘"

    def test_a_broken_listing_does_not_stop_the_import(self) -> None:
        found = _listing(
            safe_entries(
                archive({"tessera-metadata.json": "{это не json", "Своя.md": "текст"})
            )
        )
        assert found.known == {}


class TestParents:
    def test_a_page_at_the_root_has_no_parent(self) -> None:
        assert _parent_for("", {}) is None

    def test_the_folder_page_becomes_the_parent(self) -> None:
        known = uuid.uuid4()
        assert _parent_for("Родитель", {"Родитель": known}) == known

    def test_the_nearest_ancestor_wins(self) -> None:
        """Промежуточный уровень бывает без собственной страницы."""
        top = uuid.uuid4()
        assert _parent_for("Родитель/Пустой", {"Родитель": top}) == top

    def test_an_unknown_folder_leaves_the_page_at_the_root(self) -> None:
        assert _parent_for("Неизвестный", {"Другой": uuid.uuid4()}) is None


class TestSources:
    def test_every_accepted_source_has_a_parser(self) -> None:
        """Принятый вид архива должен разбираться, а не только приниматься.

        Значение без разбора обещало бы перенос, которого не происходит:
        человек выбрал бы «Confluence», а получил бы разбор своей выгрузки.
        """
        assert SOURCES == ("generic", "notion", "confluence")


@needs_database
class TestSingleFile:
    async def test_a_file_becomes_a_page(
        self, session: AsyncSession, workspace, owner, space, realtime: RealtimeDouble
    ) -> None:
        service = ImportService(session, content_client(), realtime=realtime)
        page = await service.import_file(
            file_name="Отчёт.md",
            data="# Годовой отчёт\n\nтекст".encode(),
            user_id=owner.id,
            workspace_id=workspace.id,
            space_id=space.id,
        )
        assert page.title == "Годовой отчёт"
        assert page.space_id == space.id

    async def test_a_pdf_without_a_text_layer_is_refused(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Пустая страница вместо документа выглядела бы успешным ввозом."""
        service = ImportService(session, content_client())
        with pytest.raises(AppError) as error:
            await service.import_file(
                file_name="скан.pdf",
                data=_pdf_without_text(),
                user_id=owner.id,
                workspace_id=workspace.id,
                space_id=space.id,
            )
        assert error.value.code == "error.import.no_text_layer"

    async def test_a_stranger_cannot_import(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Право то же, что у обычного создания страницы."""
        stranger = uuid.uuid4()
        with pytest.raises(AppError) as error:
            await ImportService(session, content_client()).import_file(
                file_name="чужая.md",
                data=b"# Chuzhaya",
                user_id=stranger,
                workspace_id=workspace.id,
                space_id=space.id,
            )
        assert error.value.code == "error.space.access_denied"


@needs_database
class TestTableImport:
    """Ввоз таблицы.

    По умолчанию таблица становится страницей с таблицей, и это осознанно:
    превращение страницы в базу в продукте уже есть и делается одним действием.
    Базой она ввозится по просьбе — типы столбцов при этом угадываются, и
    человеку, которому нужен документ, база досталась бы против его желания.
    """

    async def test_a_table_becomes_a_page_by_default(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        page = await ImportService(session, content_client()).import_file(
            file_name="Задачи.csv",
            data="Задача,Готово\nПервая,да\n".encode(),
            user_id=owner.id,
            workspace_id=workspace.id,
            space_id=space.id,
        )
        await session.refresh(page)
        assert page.is_base is not True

    async def test_a_table_becomes_a_base_when_asked(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        page = await ImportService(session, content_client()).import_file(
            file_name="Задачи.csv",
            data="Задача,Готово,Срок\nПервая,да,2026-01-01\nВторая,нет,2026-02-01\n".encode(),
            user_id=owner.id,
            workspace_id=workspace.id,
            space_id=space.id,
            as_base=True,
        )
        await session.refresh(page)
        assert page.is_base is True
        assert page.title == "Задачи"

        properties = (
            (
                await session.execute(
                    select(BaseProperty)
                    .where(BaseProperty.page_id == page.id)
                    .order_by(BaseProperty.position.asc())
                )
            )
            .scalars()
            .all()
        )
        assert [one.name for one in properties] == ["Задача", "Готово", "Срок"]
        # Первый столбец — название строки: без свойства-названия база не
        # открывается вовсе.
        assert properties[0].type == "title"
        assert properties[0].is_primary is True
        assert properties[1].type == "checkbox"
        assert properties[2].type == "date"

        rows = (
            (
                await session.execute(
                    select(BaseRow)
                    .where(BaseRow.page_id == page.id)
                    .order_by(BaseRow.position.asc())
                )
            )
            .scalars()
            .all()
        )
        assert len(rows) == 2
        assert rows[0].cells[properties[0].id] == "Первая"
        assert rows[0].cells[properties[1].id] is True
        assert rows[1].cells[properties[1].id] is False

        views = (
            (await session.execute(select(BaseView).where(BaseView.page_id == page.id)))
            .scalars()
            .all()
        )
        # Без представления база открывается пустой: показывать строки нечем.
        assert len(views) == 1

    async def test_a_table_of_only_a_header_is_refused(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """База из одной шапки вышла бы пустой, а причина понятнее пустого экрана."""
        with pytest.raises(AppError) as failure:
            await ImportService(session, content_client()).import_file(
                file_name="Пустая.csv",
                data="Задача,Готово\n".encode(),
                user_id=owner.id,
                workspace_id=workspace.id,
                space_id=space.id,
                as_base=True,
            )
        assert failure.value.code == "error.import.no_text"

    async def test_a_document_is_never_a_base(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Просьба относится к таблицам: строк с одинаковым набором полей в
        документе нет, и база из него вышла бы бессмысленной."""
        page = await ImportService(session, content_client()).import_file(
            file_name="Заметка.md",
            data="# Заметка\n\nТекст\n".encode(),
            user_id=owner.id,
            workspace_id=workspace.id,
            space_id=space.id,
            as_base=True,
        )
        await session.refresh(page)
        assert page.is_base is not True

    async def test_an_xlsx_is_accepted(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Книга разбирается своими средствами: новой зависимости у ввоза нет."""
        sheet = (
            '<sheetData>'
            '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c></row>'
            '<row r="2"><c r="A2" t="s"><v>2</v></c><c r="B2"><v>7</v></c></row>'
            "</sheetData>"
        )
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as book:
            book.writestr(
                "xl/worksheets/sheet1.xml",
                f'<worksheet xmlns="{namespace}">{sheet}</worksheet>',
            )
            book.writestr(
                "xl/sharedStrings.xml",
                f'<sst xmlns="{namespace}">'
                "<si><t>Задача</t></si><si><t>Часы</t></si><si><t>Первая</t></si>"
                "</sst>",
            )

        page = await ImportService(session, content_client()).import_file(
            file_name="План.xlsx",
            data=buffer.getvalue(),
            user_id=owner.id,
            workspace_id=workspace.id,
            space_id=space.id,
            as_base=True,
        )
        await session.refresh(page)
        assert page.is_base is True

        properties = (
            (
                await session.execute(
                    select(BaseProperty)
                    .where(BaseProperty.page_id == page.id)
                    .order_by(BaseProperty.position.asc())
                )
            )
            .scalars()
            .all()
        )
        assert [one.name for one in properties] == ["Задача", "Часы"]
        assert properties[1].type == "number"


class TestArchiveTask:
    async def test_the_tree_follows_the_folders(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        task = await _task(session, workspace, owner, space)
        created = await ImportService(session, content_client())._unpack(
            task,
            archive(
                {
                    "Родитель.md": "# Родитель",
                    "Родитель/Потомок.md": "# Потомок",
                }
            ),
        )
        assert created == 2

        pages = await _pages_of(session, space.id, ("Родитель", "Потомок"))
        assert pages["Потомок"].parent_page_id == pages["Родитель"].id

    async def test_the_icon_from_the_listing_is_applied(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        task = await _task(session, workspace, owner, space)
        meta = {"pages": {"Со значком.md": {"icon": "📘"}}}
        await ImportService(session, content_client())._unpack(
            task,
            archive(
                {
                    "tessera-metadata.json": json.dumps(meta),
                    "Со значком.md": "# Со значком",
                }
            ),
        )
        pages = await _pages_of(session, space.id, ("Со значком",))
        assert pages["Со значком"].icon == "📘"

    async def test_the_base_from_the_listing_is_restored(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Содержимое базы приходит оглавлением, а не документом.

        В самом файле у базы один заголовок: строки и свойства в документе не
        живут. Без восстановления ввезённая база оставалась пустой страницей.
        """
        task = await _task(session, workspace, owner, space)
        meta = {
            "pages": {
                "Задачи.md": {
                    "base": {
                        "schemaVersion": 1,
                        "properties": [
                            {
                                "id": "p1",
                                "name": "Название",
                                "type": "title",
                                "position": "h0",
                                "isPrimary": True,
                                "typeOptions": None,
                            }
                        ],
                        "views": [
                            {
                                "name": "Таблица",
                                "type": "table",
                                "position": "h0",
                                "config": {"groupByPropertyId": "p1"},
                            }
                        ],
                        "rows": [{"cells": {"p1": "Первая"}, "position": "h0"}],
                    }
                }
            }
        }
        await ImportService(session, content_client())._unpack(
            task,
            archive({"tessera-metadata.json": json.dumps(meta), "Задачи.md": "# Задачи"}),
        )

        pages = await _pages_of(session, space.id, ("Задачи",))
        page = pages["Задачи"]
        await session.refresh(page)
        assert page.is_base is True

        properties = (
            (
                await session.execute(
                    select(BaseProperty).where(BaseProperty.page_id == page.id)
                )
            )
            .scalars()
            .all()
        )
        # Идентификатор переносится как есть: на него ссылаются и ячейки, и
        # настройки представления.
        assert [one.id for one in properties] == ["p1"]

        views = (
            (await session.execute(select(BaseView).where(BaseView.page_id == page.id)))
            .scalars()
            .all()
        )
        assert views[0].config == {"groupByPropertyId": "p1"}

        rows = (
            (await session.execute(select(BaseRow).where(BaseRow.page_id == page.id)))
            .scalars()
            .all()
        )
        assert [one.cells for one in rows] == [{"p1": "Первая"}]

    async def test_the_embedded_base_points_at_the_imported_page(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Встроенная база после ввоза находит свою страницу.

        Узел встроенной базы держит идентификатор страницы, а не адрес. При
        ввозе идентификаторы новые, и без подстановки на месте таблицы
        показывалось «база не найдена» — при том, что сама база ввезена рядом
        и целиком. Найдено глазами на стенде после оборота «вывоз — ввоз».
        """
        was = "5ae4c423-5ca7-4779-886b-a21f1a260028"

        def handler(request: httpx.Request) -> httpx.Response:
            return httpx.Response(
                200,
                json={
                    "content": {
                        "type": "doc",
                        "content": [{"type": "base", "attrs": {"pageId": was}}],
                    }
                },
            )

        client = ContentClient("http://collab:3001", transport=httpx.MockTransport(handler))
        task = await _task(session, workspace, owner, space)
        meta = {
            "pages": {
                "Задачи.md": {
                    "pageId": was,
                    "base": {
                        "schemaVersion": 1,
                        "properties": [
                            {
                                "id": "p1",
                                "name": "Название",
                                "type": "title",
                                "position": "h0",
                                "isPrimary": True,
                                "typeOptions": None,
                            }
                        ],
                        "views": [],
                        "rows": [],
                    },
                },
                "Страница.md": {"pageId": "9d3a6c1e-0000-4000-8000-000000000001"},
            }
        }
        await ImportService(session, client)._unpack(
            task,
            archive(
                {
                    "tessera-metadata.json": json.dumps(meta),
                    "Задачи.md": "# Задачи",
                    "Страница.md": "# Страница",
                }
            ),
        )

        pages = await _pages_of(session, space.id, ("Задачи", "Страница"))
        embedded = pages["Страница"].content["content"][0]["attrs"]["pageId"]
        assert embedded == str(pages["Задачи"].id)

    async def test_an_unknown_embedded_base_is_left_alone(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Чужой идентификатор остаётся как есть.

        Базы этой выгрузки в архиве нет, подставлять нечего. Угадывание по
        названию поставило бы на её место чужую базу — это хуже, чем честное
        «база не найдена».
        """
        was = "5ae4c423-5ca7-4779-886b-a21f1a260028"

        def handler(request: httpx.Request) -> httpx.Response:
            return httpx.Response(
                200,
                json={
                    "content": {
                        "type": "doc",
                        "content": [{"type": "base", "attrs": {"pageId": was}}],
                    }
                },
            )

        client = ContentClient("http://collab:3001", transport=httpx.MockTransport(handler))
        task = await _task(session, workspace, owner, space)
        meta = {"pages": {"Страница.md": {"pageId": "9d3a6c1e-0000-4000-8000-000000000001"}}}
        await ImportService(session, client)._unpack(
            task,
            archive(
                {"tessera-metadata.json": json.dumps(meta), "Страница.md": "# Страница"}
            ),
        )

        pages = await _pages_of(session, space.id, ("Страница",))
        assert pages["Страница"].content["content"][0]["attrs"]["pageId"] == was

    async def test_the_context_from_the_listing_is_restored(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Обсуждение, метки и проверка приходят оглавлением, а не документом.

        Ни markdown, ни HTML их не несут: они живут рядом со страницей. Без
        восстановления полная выгрузка теряла бы всё, кроме текста.
        """
        task = await _task(session, workspace, owner, space)
        name = f"метка-{uuid.uuid4().hex[:6]}"
        meta = {
            "pages": {
                "Страница.md": {
                    "comments": [
                        {
                            "content": {
                                "type": "doc",
                                "content": [
                                    {
                                        "type": "paragraph",
                                        "content": [{"type": "text", "text": "Первая"}],
                                    }
                                ],
                            },
                            "authorEmail": owner.email,
                            "createdAt": "2026-01-01T00:00:00+00:00",
                        },
                        {
                            "content": {
                                "type": "doc",
                                "content": [
                                    {
                                        "type": "paragraph",
                                        "content": [{"type": "text", "text": "Ответ"}],
                                    }
                                ],
                            },
                            "parentIndex": 0,
                            "authorEmail": "нет-такого@example.com",
                        },
                    ],
                    "labels": [name],
                    "verification": {
                        "type": "manual",
                        "status": "verified",
                        "verifiedAt": "2026-02-01T00:00:00+00:00",
                        "verifierEmails": [owner.email],
                    },
                }
            }
        }
        await ImportService(session, content_client())._unpack(
            task,
            archive(
                {"tessera-metadata.json": json.dumps(meta), "Страница.md": "# Страница"}
            ),
        )

        pages = await _pages_of(session, space.id, ("Страница",))
        page = pages["Страница"]

        comments = (
            (
                await session.execute(
                    select(Comment)
                    .where(Comment.page_id == page.id)
                    .order_by(Comment.created_at.asc())
                )
            )
            .scalars()
            .all()
        )
        assert len(comments) == 2
        # Автор найден по почте, а не по идентификатору: тот принадлежит
        # прежней вике.
        assert comments[0].creator_id == owner.id
        # Ветвление сохранено номером в перечне: идентификаторы после ввоза
        # другие, а порядок тот же.
        assert comments[1].parent_comment_id == comments[0].id
        # Неизвестный автор не отменяет реплику: она достаётся ввозящему.
        assert comments[1].creator_id == task.creator_id

        labels = (
            (
                await session.execute(
                    select(Label.name)
                    .join(PageLabel, PageLabel.label_id == Label.id)
                    .where(PageLabel.page_id == page.id)
                )
            )
            .scalars()
            .all()
        )
        assert list(labels) == [name]

        verification = (
            await session.execute(
                select(PageVerification).where(PageVerification.page_id == page.id)
            )
        ).scalar_one()
        assert verification.status == "verified"
        verifiers = (
            (
                await session.execute(
                    select(PageVerifier.user_id).where(
                        PageVerifier.page_verification_id == verification.id
                    )
                )
            )
            .scalars()
            .all()
        )
        assert list(verifiers) == [owner.id]

    async def test_a_mention_comes_back_as_a_mention(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Упоминание, ставшее при вывозе ссылкой, возвращается упоминанием.

        Вывоз разворачивает его намеренно: узла упоминания в чужом редакторе
        нет. Обратно оно возвращается по списку из оглавления, а не по виду
        адреса — иначе всякая ссылка на свою страницу становилась бы
        упоминанием.
        """

        def handler(request: httpx.Request) -> httpx.Response:
            body = json.loads(request.content)
            markdown = str(body.get("markdown") or body.get("html") or "")
            # Ссылка на соседний файл архива — то, во что вывоз превратил
            # упоминание.
            return httpx.Response(
                200,
                json={
                    "content": {
                        "type": "doc",
                        "content": [
                            {
                                "type": "paragraph",
                                "content": [
                                    {
                                        "type": "text",
                                        "text": "Соседняя",
                                        "marks": [
                                            {
                                                "type": "link",
                                                "attrs": {"href": "Соседняя.md"},
                                            }
                                        ],
                                    }
                                ],
                            }
                        ],
                    }
                    if "Ссылающаяся" in markdown
                    else {"type": "doc", "content": [{"type": "paragraph"}]}
                },
            )

        client = ContentClient("http://collab:3001", transport=httpx.MockTransport(handler))
        task = await _task(session, workspace, owner, space)
        meta = {
            "pages": {
                "Ссылающаяся.md": {"mentions": [{"href": "Соседняя.md", "label": "Соседняя"}]},
                "Соседняя.md": {},
            }
        }
        await ImportService(session, client)._unpack(
            task,
            archive(
                {
                    "tessera-metadata.json": json.dumps(meta),
                    "Ссылающаяся.md": "# Ссылающаяся",
                    "Соседняя.md": "# Соседняя",
                }
            ),
        )

        pages = await _pages_of(session, space.id, ("Ссылающаяся", "Соседняя"))
        node = pages["Ссылающаяся"].content["content"][0]["content"][0]
        assert node["type"] == "mention"
        assert node["attrs"]["entityType"] == "page"
        assert node["attrs"]["entityId"] == str(pages["Соседняя"].id)
        assert node["attrs"]["slugId"] == pages["Соседняя"].slug_id

    async def test_an_ordinary_link_stays_a_link(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Ссылка, которой не было в списке, остаётся ссылкой.

        На одну и ту же страницу в тексте бывает и упоминание, и обычная
        ссылка: превращать в упоминание обе значило бы решать за человека.
        """

        def handler(request: httpx.Request) -> httpx.Response:
            body = json.loads(request.content)
            markdown = str(body.get("markdown") or body.get("html") or "")
            return httpx.Response(
                200,
                json={
                    "content": {
                        "type": "doc",
                        "content": [
                            {
                                "type": "paragraph",
                                "content": [
                                    {
                                        "type": "text",
                                        "text": "смотри тут",
                                        "marks": [
                                            {
                                                "type": "link",
                                                "attrs": {"href": "Соседняя.md"},
                                            }
                                        ],
                                    }
                                ],
                            }
                        ],
                    }
                    if "Ссылающаяся" in markdown
                    else {"type": "doc", "content": [{"type": "paragraph"}]}
                },
            )

        client = ContentClient("http://collab:3001", transport=httpx.MockTransport(handler))
        task = await _task(session, workspace, owner, space)
        meta = {
            "pages": {
                # Подпись в списке другая: это упоминание, а не эта ссылка.
                "Ссылающаяся.md": {"mentions": [{"href": "Соседняя.md", "label": "Соседняя"}]},
                "Соседняя.md": {},
            }
        }
        await ImportService(session, client)._unpack(
            task,
            archive(
                {
                    "tessera-metadata.json": json.dumps(meta),
                    "Ссылающаяся.md": "# Ссылающаяся",
                    "Соседняя.md": "# Соседняя",
                }
            ),
        )

        pages = await _pages_of(session, space.id, ("Ссылающаяся", "Соседняя"))
        node = pages["Ссылающаяся"].content["content"][0]["content"][0]
        assert node["type"] == "text"
        # Адрес при этом подставлен настоящий: ссылка обязана работать.
        assert node["marks"][0]["attrs"]["href"].endswith(pages["Соседняя"].slug_id)

    async def test_a_share_is_restored_with_a_new_key(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Ключ заводится свой.

        Ключ и есть учётные данные того, кто открывает страницу без входа:
        перенос ключа раздавал бы доступ вместе с файлом архива.
        """
        task = await _task(session, workspace, owner, space)
        meta = {
            "pages": {
                "Открытая.md": {
                    "share": {"includeSubPages": True, "searchIndexing": False},
                }
            }
        }
        await ImportService(session, content_client())._unpack(
            task,
            archive(
                {"tessera-metadata.json": json.dumps(meta), "Открытая.md": "# Открытая"}
            ),
        )

        pages = await _pages_of(session, space.id, ("Открытая",))
        share = (
            await session.execute(
                select(Share).where(Share.page_id == pages["Открытая"].id)
            )
        ).scalar_one()
        assert share.key
        assert bool(share.include_sub_pages) is True

    async def test_a_foreign_archive_restores_nothing_extra(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Чужая выгрузка снимка не несёт, и шаг не делает ничего."""
        task = await _task(session, workspace, owner, space)
        await ImportService(session, content_client())._unpack(
            task, archive({"Страница.md": "# Страница"})
        )

        pages = await _pages_of(session, space.id, ("Страница",))
        found = (
            await session.execute(
                select(func.count())
                .select_from(Comment)
                .where(Comment.page_id == pages["Страница"].id)
            )
        ).scalar_one()
        assert found == 0

    async def test_a_base_without_properties_stays_an_ordinary_page(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """База без единого свойства не открывается: показывать нечего."""
        task = await _task(session, workspace, owner, space)
        meta = {"pages": {"Задачи.md": {"base": {"properties": [], "views": [], "rows": []}}}}
        await ImportService(session, content_client())._unpack(
            task,
            archive({"tessera-metadata.json": json.dumps(meta), "Задачи.md": "# Задачи"}),
        )

        pages = await _pages_of(session, space.id, ("Задачи",))
        await session.refresh(pages["Задачи"])
        assert pages["Задачи"].is_base is False

    async def test_an_unreadable_entry_does_not_cancel_the_archive(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Один битый файл из тысячи не повод отменить весь ввоз."""
        task = await _task(session, workspace, owner, space)
        created = await ImportService(session, content_client())._unpack(
            task,
            archive({"Целая.md": "# Целая", "Битая.pdf": "не pdf".encode()}),
        )
        assert created == 1

    async def test_an_archive_without_documents_is_refused(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        task = await _task(session, workspace, owner, space)
        with pytest.raises(AppError) as error:
            await ImportService(session, content_client())._unpack(
                task, archive({"картинка.png": b"\x89PNG"})
            )
        assert error.value.code == "error.import.nothing_to_import"

    async def test_a_failure_is_written_into_the_task(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Человек опрашивает задание и без причины приходит спрашивать."""
        task = await _task(session, workspace, owner, space)
        storage = _StorageDouble(archive({"картинка.png": b"\x89PNG"}))
        service = ImportService(session, content_client(), storage=storage)

        assert await service.run_archive(task.id) == 0
        await session.refresh(task)
        assert task.status == STATUS_FAILED
        assert task.error_message == "error.import.nothing_to_import"

    async def test_the_uploaded_file_survives_a_failure(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """По оставленному файлу разбираются, что пошло не так."""
        task = await _task(session, workspace, owner, space)
        storage = _StorageDouble(archive({"картинка.png": b"\x89PNG"}))
        await ImportService(session, content_client(), storage=storage).run_archive(task.id)
        assert storage.deleted == []

    async def test_success_removes_the_uploaded_file(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        task = await _task(session, workspace, owner, space)
        storage = _StorageDouble(archive({"Страница.md": "# Страница"}))
        service = ImportService(session, content_client(), storage=storage)

        assert await service.run_archive(task.id) == 1
        await session.refresh(task)
        assert task.status == STATUS_SUCCESS
        assert storage.deleted == [task.file_path]

    async def test_a_finished_task_is_not_run_twice(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Повтор задания не должен завозить те же страницы второй раз."""
        task = await _task(session, workspace, owner, space, status=STATUS_SUCCESS)
        storage = _StorageDouble(archive({"Страница.md": "# Страница"}))
        assert await ImportService(session, content_client(), storage=storage).run_archive(
            task.id
        ) == 0
        assert storage.read == []

    async def test_a_deployment_without_storage_records_the_reason(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Иначе задание падает разыменованием пустоты и причина теряется."""
        task = await _task(session, workspace, owner, space)
        assert await ImportService(session, content_client()).run_archive(task.id) == 0
        await session.refresh(task)
        assert task.error_message == "error.import.unavailable"

    async def test_a_missing_task_is_not_a_crash(self, session: AsyncSession) -> None:
        service = ImportService(session, content_client(), storage=_StorageDouble(b""))
        assert await service.run_archive(uuid.uuid4()) == 0

    async def test_a_stranger_cannot_schedule_an_archive(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        service = ImportService(
            session,
            content_client(),
            storage=_StorageDouble(b""),
            queue=_QueueDouble(),
        )
        with pytest.raises(AppError) as error:
            await service.schedule_archive(
                file_name="выгрузка.zip",
                data=archive({"Страница.md": "# Страница"}),
                source="generic",
                user_id=uuid.uuid4(),
                workspace_id=workspace.id,
                space_id=space.id,
            )
        assert error.value.code == "error.space.space_not_found"

    async def test_a_reader_cannot_schedule_an_archive(
        self, session: AsyncSession, workspace, space
    ) -> None:
        """Архив заводит страницы, и читателю пространства этого нельзя.

        Проверка в самой службе, а не только в маршруте: задание переживает
        запрос, и представление о правах у него своего быть не должно.
        """
        reader = await _member(session, workspace, space, role="reader")
        service = ImportService(
            session,
            content_client(),
            storage=_StorageDouble(b""),
            queue=_QueueDouble(),
        )
        with pytest.raises(AppError) as error:
            await service.schedule_archive(
                file_name="выгрузка.zip",
                data=archive({"Страница.md": "# Страница"}),
                source="generic",
                user_id=reader,
                workspace_id=workspace.id,
                space_id=space.id,
            )
        assert error.value.code == "error.space.access_denied"

    async def test_a_scheduled_archive_waits_for_the_worker(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Страницы заводит исполнитель, а не запрос: их тысячи."""
        storage = _StorageDouble(b"")
        queue = _QueueDouble()
        task = await ImportService(
            session, content_client(), storage=storage, queue=queue
        ).schedule_archive(
            file_name="выгрузка.zip",
            data=archive({"Страница.md": "# Страница"}),
            source="generic",
            user_id=owner.id,
            workspace_id=workspace.id,
            space_id=space.id,
        )

        assert task.status == STATUS_PROCESSING
        assert queue.jobs == [("import-task", {"task_id": str(task.id)})]
        assert storage.put_keys and str(task.id) in storage.put_keys[0]


class _StorageDouble:
    """Хранилище в памяти. Настоящее здесь проверяло бы MinIO, а не ввоз."""

    def __init__(self, data: bytes) -> None:
        self._data = data
        self.deleted: list[str] = []
        self.read: list[str] = []
        self.put_keys: list[str] = []

    async def put(self, key: str, data: bytes, content_type: str | None = None) -> None:
        self.put_keys.append(key)
        self._data = data

    async def get(self, key: str) -> bytes:
        self.read.append(key)
        return self._data

    async def delete(self, key: str) -> None:
        self.deleted.append(key)


class _QueueDouble:
    def __init__(self) -> None:
        self.jobs: list[tuple[str, dict]] = []

    async def enqueue(self, name: str, *args: object, **payload: object) -> bool:
        self.jobs.append((name, payload))
        return True


async def _task(
    session: AsyncSession,
    workspace,
    owner,
    space,
    *,
    status: str = STATUS_PROCESSING,
    source: str = "generic",
) -> FileTask:
    task_id = uuid.uuid4()
    await session.execute(
        insert(FileTask).values(
            id=task_id,
            type="import",
            source=source,
            status=status,
            file_name="выгрузка.zip",
            file_path=f"{workspace.id}/imports/{task_id}/выгрузка.zip",
            file_size=1,
            file_ext=".zip",
            creator_id=owner.id,
            space_id=space.id,
            workspace_id=workspace.id,
        )
    )
    await session.commit()
    return await session.get(FileTask, task_id)


async def _member(session: AsyncSession, workspace, space, *, role: str) -> uuid.UUID:
    """Участник пространства с заданной ролью."""
    user_id = uuid.uuid4()
    await session.execute(
        insert(User).values(
            id=user_id,
            name="Участник",
            email=f"{user_id}@example.org",
            password="x",
            role="member",
            workspace_id=workspace.id,
        )
    )
    await session.execute(
        insert(SpaceMember).values(
            id=uuid.uuid4(),
            user_id=user_id,
            space_id=space.id,
            role=role,
            added_by_id=user_id,
        )
    )
    await session.commit()
    return user_id


async def _pages_of(session: AsyncSession, space_id: uuid.UUID, titles: tuple[str, ...]) -> dict:
    found = (
        (
            await session.execute(
                select(Page)
                .where(Page.space_id == space_id)
                .where(Page.title.in_(titles))
                .where(Page.deleted_at.is_(None))
            )
        )
        .scalars()
        .all()
    )
    return {one.title: one for one in found}


def _pdf_without_text() -> bytes:
    """Наименьший PDF с одной пустой страницей.

    Собран здесь, а не взят файлом: разбор обязан отличать «текста нет» от
    «файл не разобрался», и для этого нужен настоящий PDF без текста.
    """
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    start = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{start}\n%%EOF\n"
    ).encode()
    return bytes(out)


@needs_database
class TestMembership:
    async def test_the_space_member_fixture_is_real(
        self, session: AsyncSession, owner, space
    ) -> None:
        """Опора остальных проверок: владелец действительно состоит в пространстве."""
        found = (
            await session.execute(
                select(SpaceMember)
                .where(SpaceMember.space_id == space.id)
                .where(SpaceMember.user_id == owner.id)
                .where(SpaceMember.deleted_at.is_(None))
            )
        ).scalars().first()
        assert found is not None


CONFLUENCE_INDEX = """<html><body><div id="main-content"><div class="pageSection">
  <ul>
    <li>
      <a href="Reglament_1.html">Регламент</a>
      <ul><li><a href="Prilozhenie_2.html">Приложение</a></li></ul>
    </li>
  </ul>
</div></div></body></html>"""


def _confluence_page(title: str, body: str) -> str:
    return f"""<html><body>
      <div id="breadcrumb-section"><ol class="breadcrumb"><li>Пространство</li></ol></div>
      <h1><span id="title-text">{title}</span></h1>
      <div id="main-content">
        <div class="page-metadata">Создано пользователем Иванов</div>
        <p>{body}</p>
        <div class="pageSection group"><h2>Attachments:</h2></div>
      </div>
      <div id="footer">Экспортировано Confluence</div>
    </body></html>"""


#: Названия страниц проверочных выгрузок. Приметные нарочно: база проверок
#: общая с рабочей, и обычное имя нашлось бы в ней и без ввоза.
CONF_TOP = "Регламент выгрузки Confluence"
CONF_CHILD = "Приложение выгрузки Confluence"
NOTION_TOP = "Регламент выгрузки Notion"
NOTION_CHILD = "Приложение выгрузки Notion"
NOTION_ID_ONE = "1f2e3d4c5b6a7980a1b2c3d4e5f60718"
NOTION_ID_TWO = "0011223344556677889900aabbccddee"

CONFLUENCE_INDEX = f"""<html><body><div id="main-content"><div class="pageSection">
  <ul>
    <li>
      <a href="Reglament_1.html">{CONF_TOP}</a>
      <ul><li><a href="Prilozhenie_2.html">{CONF_CHILD}</a></li></ul>
    </li>
  </ul>
</div></div></body></html>"""


def _confluence_page(title: str, body: str) -> str:
    return f"""<html><body>
      <div id="breadcrumb-section"><ol class="breadcrumb"><li>Пространство</li></ol></div>
      <h1><span id="title-text">{title}</span></h1>
      <div id="main-content">
        <div class="page-metadata">Создано пользователем Иванов</div>
        <p>{body}</p>
        <div class="pageSection group"><h2>Attachments:</h2></div>
      </div>
      <div id="footer">Экспортировано Confluence</div>
    </body></html>"""


#: Страница с вложением. Confluence Server кладёт файл под числовым именем без
#: расширения, а настоящее имя и тип пишет рядом со ссылкой.
CONFLUENCE_WITH_FILE = """<html><body>
  <h1><span id="title-text">{title}</span></h1>
  <div id="main-content">
    <p>Смотри <a href="attachments/65601/65602">схему</a>.</p>
    <div class="pageSection group">
      <h2>Attachments:</h2>
      <div class="greybox">
        <a href="attachments/65601/65602">схема.txt</a> (text/plain)
      </div>
    </div>
  </div>
</body></html>"""


def _confluence_archive(*, with_file: bool = False) -> bytes:
    files = {
        "index.html": CONFLUENCE_INDEX,
        "Reglament_1.html": (
            CONFLUENCE_WITH_FILE.format(title=CONF_TOP)
            if with_file
            else _confluence_page(CONF_TOP, "Дамп ежедневно")
        ),
        "Prilozhenie_2.html": _confluence_page(CONF_CHILD, "Список серверов"),
    }
    if with_file:
        files["attachments/65601/65602"] = "содержимое схемы"
    return archive(files)


def _notion_archive() -> bytes:
    """Выгрузка Notion.

    У вложенной страницы заголовка в разметке нет нарочно: тогда он берётся из
    имени файла, и видно, срезан ли идентификатор. С заголовком в разметке
    проверка проходила бы и без срезания.
    """
    return archive(
        {
            f"{NOTION_TOP} {NOTION_ID_ONE}.md": f"# {NOTION_TOP}\n\nтекст",
            f"{NOTION_TOP} {NOTION_ID_ONE}/{NOTION_CHILD} {NOTION_ID_TWO}.md": "просто текст",
        }
    )


@needs_database
class TestForeignArchives:
    """Ввоз чужих выгрузок.

    Устройство архива разбирается отдельно (`test_import_archives.py`), здесь —
    что разобранное доезжает до страниц: дерево, заголовки, обвязка.

    Названия нарочно приметные: база проверок общая с рабочей, и обычное имя
    вроде «Регламент» нашлось бы и без ввоза.
    """

    async def test_a_confluence_export_keeps_its_tree(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Иерархия берётся из оглавления, а не из каталогов.

        Все страницы выгрузки лежат в корне архива плоско: по каталогам дерево
        не построить, и без оглавления они приехали бы вровень.
        """
        task = await _task(session, workspace, owner, space, source="confluence")
        storage = _StorageDouble(_confluence_archive())

        created = await ImportService(
            session, content_client(), storage=storage
        ).run_archive(task.id)

        assert created == 2
        pages = await _pages_of(session, space.id, (CONF_TOP, CONF_CHILD))
        assert pages[CONF_CHILD].parent_page_id == pages[CONF_TOP].id
        assert pages[CONF_TOP].parent_page_id is None

    async def test_a_confluence_page_arrives_without_the_wrapper(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Обвязка самой Confluence в страницу не переносится.

        Проверяется то, что уходит на преобразование: сама схема узлов живёт в
        соседней службе, и здесь она подменена — смотреть надо на вход, а не на
        выход.
        """
        task = await _task(session, workspace, owner, space, source="confluence")
        storage = _StorageDouble(_confluence_archive())
        seen: list = []

        await ImportService(session, content_client(seen), storage=storage).run_archive(task.id)

        sent = " ".join(json.dumps(one, ensure_ascii=False) for one in seen)
        assert "Дамп ежедневно" in sent
        assert "Создано пользователем" not in sent
        assert "Экспортировано Confluence" not in sent
        assert "Attachments" not in sent

    async def test_a_confluence_page_does_not_repeat_its_title(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Название выгрузка хранит отдельно от тела, но пишет и первым заголовком.

        Оставленный, он повторяет название страницы: одно и то же видно и в
        дереве, и первой строкой тела. Одиночный ввоз это снимал, архивный нет.
        """
        task = await _task(session, workspace, owner, space, source="confluence")
        storage = _StorageDouble(_confluence_archive())
        doc = {
            "type": "doc",
            "content": [
                {
                    "type": "heading",
                    "attrs": {"level": 1},
                    "content": [{"type": "text", "text": CONF_TOP}],
                },
                {"type": "paragraph", "content": [{"type": "text", "text": "Дамп ежедневно"}]},
            ],
        }

        await ImportService(session, content_client(doc=doc), storage=storage).run_archive(task.id)

        pages = await _pages_of(session, space.id, (CONF_TOP, CONF_CHILD))
        for page in pages.values():
            await session.refresh(page)
        assert [one["type"] for one in pages[CONF_TOP].content["content"]] == ["paragraph"]
        # У вложенной страницы название другое, и тот же заголовок для неё —
        # часть документа. Снимать его нельзя.
        assert [one["type"] for one in pages[CONF_CHILD].content["content"]] == [
            "heading",
            "paragraph",
        ]

    async def test_an_attachment_is_brought_in_and_its_link_rewritten(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Без этого ссылка `attachments/65601/65602` указывает в пустоту.

        Такого адреса на нашей стороне нет: файл лежит в архиве, а страница
        ссылается на него путём внутри выгрузки.
        """
        task = await _task(session, workspace, owner, space, source="confluence")
        storage = _StorageDouble(_confluence_archive(with_file=True))
        seen: list = []

        await ImportService(
            session, content_client(seen), storage=storage
        ).run_archive(task.id)

        page = (await _pages_of(session, space.id, (CONF_TOP,)))[CONF_TOP]
        saved = (
            (
                await session.execute(
                    select(Attachment).where(Attachment.page_id == page.id)
                )
            )
            .scalars()
            .all()
        )
        assert [one.file_name for one in saved] == ["схема.txt"]

        sent = " ".join(json.dumps(one, ensure_ascii=False) for one in seen)
        assert f"/api/files/{saved[0].id}/" in sent
        assert "attachments/65601/65602" not in sent

    async def test_an_archive_that_is_not_a_confluence_export_is_refused(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Иначе своя выгрузка, выбранная как чужая, разбиралась бы наугад."""
        task = await _task(session, workspace, owner, space, source="confluence")
        storage = _StorageDouble(archive({"Страница.md": "# Страница"}))

        created = await ImportService(
            session, content_client(), storage=storage
        ).run_archive(task.id)

        assert created == 0
        await session.refresh(task)
        assert task.error_message == "error.import.unknown_source"

    async def test_a_notion_export_loses_its_identifiers(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Notion приписывает идентификатор к каждому имени.

        Без срезания он попадает и в заголовок страницы, и в имя ветви: дерево
        строится по каталогам, а каталог называется так же.
        """
        task = await _task(session, workspace, owner, space, source="notion")
        storage = _StorageDouble(_notion_archive())

        created = await ImportService(
            session, content_client(), storage=storage
        ).run_archive(task.id)

        assert created == 2
        pages = await _pages_of(session, space.id, (NOTION_TOP, NOTION_CHILD))
        assert set(pages) == {NOTION_TOP, NOTION_CHILD}

        # И ни одной страницы с идентификатором в заголовке.
        with_id = await session.execute(
            select(Page.title).where(
                Page.space_id == space.id, Page.title.like(f"%{NOTION_ID_ONE}%")
            )
        )
        assert with_id.scalars().all() == []

    async def test_a_notion_export_keeps_its_tree(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        task = await _task(session, workspace, owner, space, source="notion")
        storage = _StorageDouble(_notion_archive())

        await ImportService(session, content_client(), storage=storage).run_archive(task.id)

        pages = await _pages_of(session, space.id, (NOTION_TOP, NOTION_CHILD))
        assert pages[NOTION_CHILD].parent_page_id == pages[NOTION_TOP].id


class TestAcceptedFormatsAreOneList:
    """Перечень принимаемых форматов существует дважды.

    Сервер отвергает по `SINGLE_FILE_EXTENSIONS`, а окно выбора файла в
    браузере фильтрует по `IMPORT_ACCEPT`. Разойтись они могут молча, и оба
    расхождения неприятны по-своему: формат, которого нет в окне выбора,
    человек просто не найдёт, а формат, которого нет на сервере, выберется и
    получит отказ уже после загрузки.

    Разбор текста здесь уместен по той же причине, что и в `test_compose_env`:
    проверяемое свойство и есть свойство двух списков, а исполнимого моста
    между Python и TypeScript нет.
    """

    ACCEPT = (
        Path(__file__).resolve().parents[3]
        / "apps"
        / "web"
        / "src"
        / "lib"
        / "features"
        / "page"
        / "services"
        / "transfer.ts"
    )

    def _from_client(self) -> tuple[str, ...]:
        text = self.ACCEPT.read_text(encoding="utf-8")
        found = re.search(r"IMPORT_ACCEPT\s*=\s*'([^']+)'", text)
        assert found, "перечень в клиенте не найден — имя или запись изменились"
        return tuple(found.group(1).split(","))

    def test_the_client_list_is_readable(self) -> None:
        """Сначала проверяется сам разбор: пустой список сравнялся бы с чем угодно."""
        assert len(self._from_client()) > 3

    def test_both_lists_are_the_same(self) -> None:
        assert self._from_client() == SINGLE_FILE_EXTENSIONS


@needs_database
class TestDocxImport:
    """Ввоз документа Word: разметка и картинки.

    До этой работы из DOCX брался голый текст, и ввезённый документ терял
    заголовки, списки, таблицу и все встроенные картинки. Ввоз при этом был
    успешен — потерю нечем было заметить.
    """

    def _document(self, *, with_picture: bool = False) -> bytes:
        import struct
        import zlib

        from docx import Document
        from docx.shared import Inches

        def png() -> bytes:
            def chunk(tag: bytes, data: bytes) -> bytes:
                body = tag + data
                return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

            rows = b"".join(b"\x00" + b"\xff\x00\x00" * 2 for _ in range(2))
            return (
                b"\x89PNG\r\n\x1a\n"
                + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0))
                + chunk(b"IDAT", zlib.compress(rows))
                + chunk(b"IEND", b"")
            )

        document = Document()
        document.add_heading("Договор поставки", level=1)
        document.add_paragraph("Обычный абзац.")
        if with_picture:
            document.add_picture(io.BytesIO(png()), width=Inches(1))
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    async def test_the_document_goes_through_html_not_plain_text(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Разметка доходит до преобразования.

        Проверяется то, что ушло соседнему сервису: сам он в проверках отвечает
        пустым документом, и смотреть на его ответ значило бы смотреть на
        подмену.
        """
        seen: list = []
        service = ImportService(session, content_client(seen))
        await service.import_file(
            file_name="договор.docx",
            data=self._document(),
            user_id=owner.id,
            workspace_id=workspace.id,
            space_id=space.id,
        )

        sent = [one for one in seen if "html" in one]
        assert sent, "разбор ушёл не как HTML"
        assert "<h1>Договор поставки</h1>" in sent[0]["html"]
        assert "<p>Обычный абзац.</p>" in sent[0]["html"]

    async def test_the_title_comes_from_the_heading(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        service = ImportService(session, content_client())
        page = await service.import_file(
            file_name="договор.docx",
            data=self._document(),
            user_id=owner.id,
            workspace_id=workspace.id,
            space_id=space.id,
        )
        assert page.title == "Договор поставки"

    def _titled(self, title: str, body: str) -> dict:
        """Разобранный документ, начатый заголовком с названием.

        Схема узлов живёт в соседней службе, и здесь она подменена: проверяется
        обращение с разобранным, а не сам разбор.
        """
        return {
            "type": "doc",
            "content": [
                {
                    "type": "heading",
                    "attrs": {"level": 1},
                    "content": [{"type": "text", "text": title}],
                },
                {"type": "paragraph", "content": [{"type": "text", "text": body}]},
            ],
        }

    async def test_the_title_does_not_stay_a_heading_inside(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Название взято из первого заголовка — внутри страницы он лишний.

        Оставленный, он повторяет название: одно и то же видно и в дереве, и
        первой строкой тела. У Markdown и HTML снятие было, у Word нет.
        """
        service = ImportService(
            session, content_client(doc=self._titled("Договор поставки", "Обычный абзац."))
        )
        page = await service.import_file(
            file_name="договор.docx",
            data=self._document(),
            user_id=owner.id,
            workspace_id=workspace.id,
            space_id=space.id,
        )

        await session.refresh(page)
        assert page.title == "Договор поставки"
        assert [one["type"] for one in page.content["content"]] == ["paragraph"]

    async def test_a_broken_document_is_refused(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Пустой разбор — не пустой документ, а не разобранный."""
        with pytest.raises(AppError) as error:
            await ImportService(session, content_client()).import_file(
                file_name="битый.docx",
                data=b"PK\x03\x04 not a document",
                user_id=owner.id,
                workspace_id=workspace.id,
                space_id=space.id,
            )
        assert error.value.code == "error.import.no_text"

    async def test_without_storage_the_text_still_arrives(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Картинки не переносятся, но документ ввозится.

        Отказ здесь стоил бы человеку всего документа ради одной картинки.
        """
        page = await ImportService(session, content_client()).import_file(
            file_name="с картинкой.docx",
            data=self._document(with_picture=True),
            user_id=owner.id,
            workspace_id=workspace.id,
            space_id=space.id,
        )
        assert page.id is not None

    async def test_an_image_is_stored_and_its_address_substituted(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Картинка кладётся вложением страницы, заготовка заменяется адресом.

        Без подстановки в содержимом остаётся `docx-image:0` — на экране это
        пустое место, и человек видит документ без печати, не понимая, почему.
        """
        storage = _StorageDouble(b"")
        # Содержимое приходит от соседнего сервиса; здесь он отвечает готовым
        # документом с картинкой, чтобы проверить именно подстановку адреса.
        def handler(request: httpx.Request) -> httpx.Response:
            return httpx.Response(
                200,
                json={
                    "content": {
                        "type": "doc",
                        "content": [
                            {"type": "image", "attrs": {"src": "docx-image:0"}},
                        ],
                    }
                },
            )

        service = ImportService(
            session,
            ContentClient("http://collab:3001", transport=httpx.MockTransport(handler)),
            storage=storage,
        )
        page = await service.import_file(
            file_name="с картинкой.docx",
            data=self._document(with_picture=True),
            user_id=owner.id,
            workspace_id=workspace.id,
            space_id=space.id,
        )

        assert storage.put_keys, "картинка не легла в хранилище"
        address = page.content["content"][0]["attrs"]["src"]
        assert address.startswith("/api/files/"), address
        assert "docx-image:" not in address


@needs_database
class TestDocxInsideArchive:
    """Документ Word внутри архива.

    Путь другой, а последствие то же: без выгрузки картинок в содержимом
    остаётся заготовка, на экране это битая картинка, а в хранилище пусто.
    Ввоз при этом успешен.
    """

    def _docx(self) -> bytes:
        import struct
        import zlib

        from docx import Document
        from docx.shared import Inches

        def png() -> bytes:
            def chunk(tag: bytes, data: bytes) -> bytes:
                body = tag + data
                return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

            rows = b"".join(b"\x00" + b"\xff\x00\x00" * 2 for _ in range(2))
            return (
                b"\x89PNG\r\n\x1a\n"
                + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0))
                + chunk(b"IDAT", zlib.compress(rows))
                + chunk(b"IEND", b"")
            )

        document = Document()
        document.add_heading("Из архива", level=1)
        document.add_picture(io.BytesIO(png()), width=Inches(1))
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    async def test_its_images_are_stored_too(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        storage = _StorageDouble(b"")

        def handler(request: httpx.Request) -> httpx.Response:
            return httpx.Response(
                200,
                json={
                    "content": {
                        "type": "doc",
                        "content": [{"type": "image", "attrs": {"src": "docx-image:0"}}],
                    }
                },
            )

        # Задание заводится напрямую, как в соседних проверках архива: очередь
        # здесь ни при чём, проверяется разбор.
        task = await _task(session, workspace, owner, space)
        storage._data = archive({"договор.docx": self._docx()})
        service = ImportService(
            session,
            ContentClient("http://collab:3001", transport=httpx.MockTransport(handler)),
            storage=storage,
        )
        assert await service.run_archive(task.id)

        page = (
            await session.execute(
                select(Page)
                .where(Page.space_id == space.id)
                .where(Page.title == "Из архива")
                .where(Page.deleted_at.is_(None))
            )
        ).scalars().one()

        assert storage.put_keys, "картинка из архива не легла в хранилище"
        address = page.content["content"][0]["attrs"]["src"]
        assert address.startswith("/api/files/"), address


def test_the_import_route_is_given_storage() -> None:
    """Маршрут ввоза одного файла обязан получать хранилище.

    Без него картинки документа Word не переносятся, и заметить это по коду
    службы нельзя: она молча пропускает их, потому что так и задумано — отказ
    ради одной картинки стоил бы всего документа. Проверяется поэтому сам
    маршрут: объявлен ли у него довод.
    """
    import inspect

    from tessera_api.api.imports import ImportController
    from tessera_api.infrastructure.storage import Storage

    # Обработчик обёрнут Litestar: доводы объявлены у самой функции, а не у
    # объекта маршрута.
    handler = ImportController.import_file
    signature = inspect.signature(getattr(handler, "fn", handler))
    storage = signature.parameters.get("storage")
    assert storage is not None, "маршрут не просит хранилище"
    assert Storage.__name__ in str(storage.annotation)


class TestCsvImport:
    """Таблица из CSV.

    Решение принято настоящей выгрузкой Notion: базы лежат в ней именно `.csv`,
    и без их разбора половина выгрузки пропадала бы молча.
    """

    def test_the_first_row_becomes_the_header(self) -> None:
        html = table_to_html(read_table(".csv", "Название,Статус\nЗадача,Готово\n".encode()))
        assert "<th>Название</th><th>Статус</th>" in html
        assert "<td>Задача</td><td>Готово</td>" in html

    def test_the_separator_is_taken_from_the_file(self) -> None:
        """Выгрузки приходят и с запятой, и с точкой с запятой.

        Файл со вторым разделителем, разобранный по первому, даёт таблицу из
        одного столбца — без единого отказа.
        """
        html = table_to_html(read_table(".csv", "Название;Статус\nЗадача;Готово\n".encode()))
        assert "<th>Название</th><th>Статус</th>" in html

    def test_cells_are_escaped(self) -> None:
        """Содержимое таблицы задаёт не наш код."""
        html = table_to_html(read_table(".csv", "Название\n<script>alert(1)</script>\n".encode()))
        assert "<script>" not in html
        assert "&lt;script&gt;" in html

    def test_a_byte_order_mark_does_not_leak_into_the_header(self) -> None:
        """Выгрузки из Excel начинаются с метки порядка байтов."""
        html = table_to_html(read_table(".csv", "﻿Название,Статус\nа,б\n".encode()))
        assert "<th>Название</th>" in html

    def test_empty_rows_are_dropped(self) -> None:
        html = table_to_html(read_table(".csv", "Название\n\nЗадача\n\n".encode()))
        assert html.count("<tr>") == 2

    def test_an_empty_file_gives_nothing(self) -> None:
        assert table_to_html(read_table(".csv", b"")) == ""
        assert table_to_html(read_table(".csv", b"   \n\n")) == ""


class TestNestedArchive:
    """Архив, внутри которого лежит один архив.

    Так Notion отдаёт крупные выгрузки: снаружи `full-note.zip`, внутри
    единственный `ExportBlock-…-Part-1.zip`.
    """

    def test_a_single_inner_archive_is_unwrapped(self) -> None:
        inner = archive({"страница.md": "# Заголовок"})
        outer = archive({"ExportBlock-abc-Part-1.zip": inner})
        found = _unwrapped(safe_entries(outer))
        assert [one.path for one in found] == ["страница.md"]

    def test_an_ordinary_archive_is_left_alone(self) -> None:
        found = _unwrapped(safe_entries(archive({"а.md": "1", "б.md": "2"})))
        assert sorted(one.path for one in found) == ["а.md", "б.md"]

    def test_several_archives_are_left_alone(self) -> None:
        """Несколько означало бы выгрузку из частей: её страницы надо связывать
        между частями, а это другая задача, и делать её молча нельзя."""
        outer = archive(
            {
                "часть-1.zip": archive({"а.md": "1"}),
                "часть-2.zip": archive({"б.md": "2"}),
            }
        )
        found = _unwrapped(safe_entries(outer))
        assert len(found) == 2

    def test_a_broken_inner_archive_does_not_raise(self) -> None:
        outer = archive({"битый.zip": b"not an archive"})
        found = _unwrapped(safe_entries(outer))
        assert [one.path for one in found] == ["битый.zip"]


class TestNotionTwins:
    """Notion кладёт каждую базу дважды: урезанную и полную."""

    def test_the_full_copy_takes_the_place_of_the_trimmed_one(self) -> None:
        entries = [
            ArchiveEntry(path="Задачи.csv", data=b"a"),
            ArchiveEntry(path="Задачи_all.csv", data=b"b"),
            ArchiveEntry(path="Страница.md", data=b"c"),
        ]
        left = _without_notion_twins(entries)

        # Путь урезанной, содержимое полной. Пометка `_all` — след выгрузки, а
        # не часть названия: иначе страница называлась бы «Задачи_all», а
        # ссылки выгрузки ведут на «Задачи.csv» и остались бы висеть.
        assert [one.path for one in left] == ["Задачи.csv", "Страница.md"]
        assert left[0].data == b"b"

    def test_a_lone_table_stays(self) -> None:
        entries = [ArchiveEntry(path="Задачи.csv", data=b"a")]
        assert [one.path for one in _without_notion_twins(entries)] == ["Задачи.csv"]


def linking_client(documents: dict[str, dict]) -> ContentClient:
    """Преобразователь, отвечающий заданным документом на заданный текст.

    Общий `content_client` отвечает пустым абзацем, и на нём не видно ни
    ссылок, ни картинок. Здесь ответ задаётся текстом записи: проверяется, что
    делает ввоз с адресами внутри готового документа, а не разбор Markdown.

    Совпадение по вхождению, а не по равенству: у Confluence на преобразование
    уходит разметка страницы целиком, и записывать её в проверку значило бы
    держать вторую копию обвязки выгрузки.
    """

    def handler(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        text = body.get("markdown") or body.get("html") or ""
        for mark, document in documents.items():
            if mark in text:
                return httpx.Response(200, json={"content": document})
        return httpx.Response(200, json={"content": DOC})

    return ContentClient("http://collab:3001", transport=httpx.MockTransport(handler))


def _link_doc(href: str) -> dict:
    return {
        "type": "doc",
        "content": [
            {
                "type": "paragraph",
                "content": [
                    {
                        "type": "text",
                        "text": "туда",
                        "marks": [{"type": "link", "attrs": {"href": href}}],
                    }
                ],
            }
        ],
    }


def _attachment_doc(url: str) -> dict:
    """Документ с приложенным файлом. Адрес у него в `url`, а не в `src`."""
    return {
        "type": "doc",
        "content": [
            {
                "type": "attachment",
                "attrs": {"url": url, "name": "note.txt", "mime": "text/plain", "size": 7},
            }
        ],
    }


def _image_doc(src: str) -> dict:
    return {
        "type": "doc",
        "content": [{"type": "image", "attrs": {"src": src, "alt": "картинка"}}],
    }


def _addresses(content: object, field: str, found: list[str]) -> None:
    if isinstance(content, dict):
        value = content.get(field)
        if isinstance(value, str):
            found.append(value)
        for one in content.values():
            _addresses(one, field, found)
    elif isinstance(content, list):
        for one in content:
            _addresses(one, field, found)


def _hrefs(page: Page) -> list[str]:
    found: list[str] = []
    _addresses(page.content, "href", found)
    return found


def _sources(page: Page) -> list[str]:
    found: list[str] = []
    _addresses(page.content, "src", found)
    return found


class TestArchiveTarget:
    """Разбор относительного адреса записи архива."""

    def test_an_encoded_name_is_decoded(self) -> None:
        # Выгрузка кодирует пробел как `%20`, тире как `%E2%80%94`. Без
        # раскодирования путь не совпадает ни с одной записью архива.
        assert (
            module._archive_target("NEWS%20%E2%80%94%203afe.md", "Папка")
            == "Папка/NEWS — 3afe.md"
        )

    def test_the_folder_of_the_page_is_the_starting_point(self) -> None:
        assert module._archive_target("Название/image.png", "") == "Название/image.png"

    def test_an_outside_address_is_left_alone(self) -> None:
        outside = (
            "https://example.com/a.png",
            "/api/files/1/a.png",
            "data:image/png;base64,x",
        )
        for source in outside:
            assert module._archive_target(source, "Папка") is None

    def test_an_anchor_and_a_query_do_not_become_part_of_the_path(self) -> None:
        assert module._archive_target("Сосед.md#раздел", "") == "Сосед.md"
        assert module._archive_target("Сосед.md?v=2", "") == "Сосед.md"

    def test_a_path_leading_outside_the_archive_is_refused(self) -> None:
        assert module._archive_target("../../etc/passwd", "Папка") is None

    def test_an_empty_address_is_not_a_path(self) -> None:
        assert module._archive_target("", "Папка") is None


@needs_database
class TestArchiveLinks:
    async def test_a_link_to_a_neighbour_becomes_a_page_address(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Ссылка выгрузки вела в никуда: такого адреса на нашей стороне нет."""
        task = await _task(session, workspace, owner, space)
        client = linking_client({"# Первая": _link_doc("%D0%92%D1%82%D0%BE%D1%80%D0%B0%D1%8F.md")})

        created = await ImportService(session, client)._unpack(
            task, archive({"Первая.md": "# Первая", "Вторая.md": "# Вторая"})
        )
        assert created == 2

        pages = await _pages_of(session, space.id, ("Первая", "Вторая"))
        assert _hrefs(pages["Первая"]) == [f"/s/{space.slug}/p/{pages['Вторая'].slug_id}"]

    async def test_a_link_to_a_missing_neighbour_is_left_alone(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Подменять на пустоту нечем: адрес хотя бы говорит, куда вела ссылка."""
        task = await _task(session, workspace, owner, space)
        client = linking_client({"# Одна": _link_doc("Пропавшая.md")})

        await ImportService(session, client)._unpack(task, archive({"Одна.md": "# Одна"}))

        pages = await _pages_of(session, space.id, ("Одна",))
        assert _hrefs(pages["Одна"]) == ["Пропавшая.md"]

    async def test_an_outside_link_survives_untouched(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        task = await _task(session, workspace, owner, space)
        client = linking_client({"# Одна": _link_doc("https://example.com/a")})

        await ImportService(session, client)._unpack(task, archive({"Одна.md": "# Одна"}))

        pages = await _pages_of(session, space.id, ("Одна",))
        assert _hrefs(pages["Одна"]) == ["https://example.com/a"]

    async def test_a_notion_link_matches_despite_the_identifier(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """У Notion идентификатор сидит и в имени файла, и в самой ссылке.

        Пути записей от него очищаются, а ссылка приходит с ним: без такой же
        чистки цели ни одна внутренняя ссылка выгрузки не совпадает.
        """
        task = await _task(session, workspace, owner, space, source="notion")
        target = "%D0%92%D1%82%D0%BE%D1%80%D0%B0%D1%8F%203afe8a7a2ce981d2a2cee370c40a3ba9.md"
        client = linking_client({"# Первая": _link_doc(target)})

        await ImportService(session, client)._unpack(
            task,
            archive(
                {
                    "Первая 1afe8a7a2ce981d2a2cee370c40a3ba9.md": "# Первая",
                    "Вторая 3afe8a7a2ce981d2a2cee370c40a3ba9.md": "# Вторая",
                }
            ),
        )

        pages = await _pages_of(session, space.id, ("Первая", "Вторая"))
        assert _hrefs(pages["Первая"]) == [f"/s/{space.slug}/p/{pages['Вторая'].slug_id}"]

    async def test_an_image_of_the_archive_is_carried_over(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Картинка выгрузки лежит соседним файлом и никуда не переносилась."""
        task = await _task(session, workspace, owner, space)
        client = linking_client({"# Одна": _image_doc("%D0%9E%D0%B4%D0%BD%D0%B0/image.png")})
        storage = _StorageDouble(b"")

        await ImportService(session, client, storage=storage)._unpack(
            task, archive({"Одна.md": "# Одна", "Одна/image.png": b"\x89PNG\r\n\x1a\n"})
        )

        pages = await _pages_of(session, space.id, ("Одна",))
        assert _sources(pages["Одна"])[0].startswith("/api/files/")

        rows = await session.execute(
            select(Attachment).where(Attachment.page_id == pages["Одна"].id)
        )
        saved = rows.scalars().all()
        assert [one.file_name for one in saved] == ["image.png"]

    async def test_an_attached_file_of_the_archive_is_carried_over(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Узел вложения держит адрес в `url`, а не в `src`.

        Пока разбор смотрел только на `src` и `href`, приложенный файл не
        завозился вовсе, а в документе оставался путь внутрь архива — ссылка
        вела в пустоту.
        """
        task = await _task(session, workspace, owner, space)
        client = linking_client({"# Одна": _attachment_doc("%D0%9E%D0%B4%D0%BD%D0%B0/note.txt")})
        storage = _StorageDouble(b"")

        await ImportService(session, client, storage=storage)._unpack(
            task, archive({"Одна.md": "# Одна", "Одна/note.txt": "заметка".encode()})
        )

        pages = await _pages_of(session, space.id, ("Одна",))
        found: list[str] = []
        _addresses(pages["Одна"].content, "url", found)
        assert found and found[0].startswith("/api/files/")

        rows = await session.execute(
            select(Attachment).where(Attachment.page_id == pages["Одна"].id)
        )
        assert [one.file_name for one in rows.scalars().all()] == ["note.txt"]

    async def test_the_attachment_id_follows_the_file(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Прежний идентификатор принадлежит чужой вики.

        На вид всё цело — картинка показывается по адресу, — но по
        идентификатору ходит замена файла на месте и правка диаграммы, которая
        перезаписывает своё вложение.
        """
        task = await _task(session, workspace, owner, space)
        stale = str(uuid.uuid4())
        doc = {
            "type": "doc",
            "content": [
                {
                    "type": "image",
                    "attrs": {"src": "%D0%9E%D0%B4%D0%BD%D0%B0/image.png", "attachmentId": stale},
                }
            ],
        }
        client = linking_client({"# Одна": doc})
        storage = _StorageDouble(b"")

        await ImportService(session, client, storage=storage)._unpack(
            task, archive({"Одна.md": "# Одна", "Одна/image.png": b"\x89PNG"})
        )

        pages = await _pages_of(session, space.id, ("Одна",))
        ids: list[str] = []
        _addresses(pages["Одна"].content, "attachmentId", ids)
        rows = await session.execute(
            select(Attachment).where(Attachment.page_id == pages["Одна"].id)
        )
        saved = rows.scalars().all()
        assert ids == [str(saved[0].id)]
        assert stale not in ids

    async def test_without_storage_the_text_still_arrives(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Отказ переноса картинки не должен стоить человеку всего документа."""
        task = await _task(session, workspace, owner, space)
        client = linking_client({"# Одна": _image_doc("Одна/image.png")})

        created = await ImportService(session, client)._unpack(
            task, archive({"Одна.md": "# Одна", "Одна/image.png": b"\x89PNG"})
        )
        assert created == 1

        pages = await _pages_of(session, space.id, ("Одна",))
        assert _sources(pages["Одна"]) == ["Одна/image.png"]


@needs_database
class TestArchiveLinksElsewhere:
    """Тот же класс в двух других местах."""

    async def test_a_link_to_a_trimmed_table_leads_to_the_full_one(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Ввоз оставляет полную копию, а ссылки в выгрузке ведут на урезанную.

        Без подмены каждая ссылка на базу Notion оставалась бы висеть.
        """
        task = await _task(session, workspace, owner, space, source="notion")
        client = linking_client({"# Одна": _link_doc("%D0%97%D0%B0%D0%B4%D0%B0%D1%87%D0%B8.csv")})

        await ImportService(session, client)._unpack(
            task,
            archive(
                {
                    "Одна.md": "# Одна",
                    "Задачи.csv": "имя,срок\nпервая,завтра\n",
                    "Задачи_all.csv": "имя,срок\nпервая,завтра\nвторая,послезавтра\n",
                }
            ),
        )

        # Страница названа «Задачи», а не «Задачи_all»: пометка полной копии
        # — след выгрузки, и полная копия заняла путь урезанной.
        pages = await _pages_of(session, space.id, ("Одна", "Задачи"))
        assert _hrefs(pages["Одна"]) == [f"/s/{space.slug}/p/{pages['Задачи'].slug_id}"]

    async def test_a_document_never_becomes_an_attachment(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Неразобранный документ вложением не кладётся.

        Иначе отказ разбора выдавался бы за успех: в содержимом появлялась бы
        ссылка на файл, который должен был стать страницей.
        """
        task = await _task(session, workspace, owner, space)
        client = linking_client({"# Одна": _link_doc("Битая.pdf")})
        storage = _StorageDouble(b"")

        await ImportService(session, client, storage=storage)._unpack(
            task, archive({"Одна.md": "# Одна", "Битая.pdf": "не pdf".encode()})
        )

        pages = await _pages_of(session, space.id, ("Одна",))
        assert _hrefs(pages["Одна"]) == ["Битая.pdf"]

        rows = await session.execute(
            select(Attachment).where(Attachment.page_id == pages["Одна"].id)
        )
        assert rows.scalars().all() == []

    async def test_a_confluence_link_between_pages_becomes_a_page_address(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """У Confluence страницы ссылаются друг на друга тем же способом."""
        task = await _task(session, workspace, owner, space, source="confluence")
        storage = _StorageDouble(_confluence_archive())
        client = linking_client({"Дамп ежедневно": _link_doc("Prilozhenie_2.html")})

        created = await ImportService(session, client, storage=storage).run_archive(task.id)
        assert created == 2

        pages = await _pages_of(session, space.id, (CONF_TOP, CONF_CHILD))
        assert _hrefs(pages[CONF_TOP]) == [f"/s/{space.slug}/p/{pages[CONF_CHILD].slug_id}"]


@needs_database
class TestPdfImport:
    """Ввоз PDF идёт через сервис преобразования.

    Разбор своим средством на Python отдавал голый текст: ввезённый документ
    терял заголовки и списки, а ввоз при этом был успешен — потерю нечем было
    заметить.
    """

    async def test_the_file_goes_to_the_transform_service(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        seen: list = []
        service = ImportService(
            session, content_client(seen, pdf="<h1>Регламент</h1><p>текст</p>")
        )
        await service.import_file(
            file_name="регламент.pdf",
            data=_pdf_without_text(),
            user_id=owner.id,
            workspace_id=workspace.id,
            space_id=space.id,
        )

        # Файл уходит в base64: у сервиса один вид тела.
        sent = next(one for one in seen if "pdf" in one)
        assert base64.b64decode(sent["pdf"]) == _pdf_without_text()

    async def test_the_title_comes_from_the_markup(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Заголовок берётся из разобранного документа, а не из имени файла."""
        service = ImportService(session, content_client(pdf="<h1>Регламент</h1>"))
        page = await service.import_file(
            file_name="untitled.pdf",
            data=_pdf_without_text(),
            user_id=owner.id,
            workspace_id=workspace.id,
            space_id=space.id,
        )
        assert page.title == "Регламент"

    def _titled(self, title: str, body: str) -> dict:
        """Разобранный документ, начатый заголовком с названием.

        Схема узлов живёт в соседней службе, и здесь она подменена: проверяется
        обращение с разобранным, а не сам разбор.
        """
        return {
            "type": "doc",
            "content": [
                {
                    "type": "heading",
                    "attrs": {"level": 1},
                    "content": [{"type": "text", "text": title}],
                },
                {"type": "paragraph", "content": [{"type": "text", "text": body}]},
            ],
        }

    async def test_the_title_does_not_stay_a_heading_inside(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Название взято из первого заголовка — внутри страницы он лишний."""
        service = ImportService(
            session,
            content_client(
                pdf="<h1>Регламент</h1><p>текст</p>", doc=self._titled("Регламент", "текст")
            ),
        )
        page = await service.import_file(
            file_name="регламент.pdf",
            data=_pdf_without_text(),
            user_id=owner.id,
            workspace_id=workspace.id,
            space_id=space.id,
        )

        await session.refresh(page)
        assert page.title == "Регламент"
        assert [one["type"] for one in page.content["content"]] == ["paragraph"]

    async def test_the_refusal_of_the_service_reaches_the_person(
        self, session: AsyncSession, workspace, owner, space
    ) -> None:
        """Причина отказа называется своим кодом, а не общим «не удалось».

        У разбора PDF причин три, и человеку они говорят разное: пустой файл,
        неразобранный файл и скан без текстового слоя.
        """
        service = ImportService(session, content_client())
        with pytest.raises(AppError) as error:
            await service.import_file(
                file_name="скан.pdf",
                data=_pdf_without_text(),
                user_id=owner.id,
                workspace_id=workspace.id,
                space_id=space.id,
            )
        assert error.value.code == "error.import.no_text_layer"


class TestTransformFailures:
    """Перевод причины отказа соседа в код приложения."""

    def test_a_named_reason_becomes_its_code(self) -> None:
        response = httpx.Response(400, json={"error": "pdf_no_text_layer"})
        assert _failure_code(response, PDF_FAILURES) == "error.import.no_text_layer"

    def test_an_unnamed_reason_stays_a_general_refusal(self) -> None:
        response = httpx.Response(400, json={"error": "чего-то новое"})
        assert _failure_code(response, PDF_FAILURES) == "error.content.transform_failed"

    def test_a_body_that_is_not_ours_stays_a_general_refusal(self) -> None:
        # Так отвечает не сервис, а что-то на его месте: обратный прокси,
        # заглушка развёртывания, чужой процесс на том же порту.
        response = httpx.Response(502, text="<html>502</html>")
        assert _failure_code(response, PDF_FAILURES) == "error.content.transform_failed"

    def test_without_a_mapping_every_refusal_is_general(self) -> None:
        response = httpx.Response(400, json={"error": "pdf_no_text_layer"})
        assert _failure_code(response, None) == "error.content.transform_failed"
